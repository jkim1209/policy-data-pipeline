"""엑셀 구조 조사. 변환하지 않고 '무엇이 어디에 있는지'만 봅니다.

내부 자료 엑셀은 사람이 보라고 만든 문서입니다. 기계가 읽는 표가 아닙니다.
다단 헤더 · 병합 셀 · 단위 혼재 · 각주가 섞여 있어, 구조를 확인하지 않고
pandas.read_excel 을 바로 부르면 헤더가 값으로 들어가거나 각주가 데이터가 됩니다.

출력
  1) 시트 목록과 사용 범위
  2) 셀 덤프 (값 · 자료형 · 병합 여부)
  3) 병합 셀 범위
  4) 헤더 단수 추정 — 위에서부터 숫자가 아닌 행이 몇 줄인지
  5) 열별 자료형 혼재 / 단위 문자열 탐지
  6) 각주 후보 — 데이터 블록 아래 텍스트 행
"""
from __future__ import annotations

import argparse
import os
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# 단위로 보이는 토큰. 값 열에 섞여 들어오면 float 변환이 터집니다.
RE_UNIT = re.compile(r"(명|원|천원|백만원|억원|%|퍼센트|건|개소|인|세|㎡|가구|천명|만원)")
RE_NUMERIC_TEXT = re.compile(r"^\s*-?[\d,]+(\.\d+)?\s*$")
RE_FOOTNOTE = re.compile(r"^\s*([*※注주]|\(?\d+\)|[가-힣]{0,3}\s*[:：])")


def cell_kind(v):
    if v is None:
        return "빈칸"
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, (int, float)):
        return "숫자"
    s = str(v)
    if RE_NUMERIC_TEXT.match(s):
        return "숫자형문자"          # 가장 위험합니다. 보기엔 숫자, 타입은 문자
    return "문자"


def merged_lookup(ws):
    """셀 좌표 → 병합 범위 문자열."""
    out = {}
    for rng in ws.merged_cells.ranges:
        for row in ws[rng.coord]:
            for c in row:
                out[c.coordinate] = str(rng)
    return out


def dump(ws, max_rows, max_cols):
    merged = merged_lookup(ws)
    print(f"\n### 시트 `{ws.title}` — {ws.max_row}행 × {ws.max_column}열")
    if ws.merged_cells.ranges:
        print(f"병합 범위 {len(ws.merged_cells.ranges)}개: "
              + ", ".join(str(r) for r in ws.merged_cells.ranges))
    else:
        print("병합 범위 없음")

    print("\n#### 셀 덤프")
    header = ["행"] + [get_column_letter(c) for c in range(1, min(ws.max_column, max_cols) + 1)]
    print("| " + " | ".join(header) + " |")
    print("|" + "---|" * len(header))
    for r in range(1, min(ws.max_row, max_rows) + 1):
        cells = []
        for c in range(1, min(ws.max_column, max_cols) + 1):
            cell = ws.cell(row=r, column=c)
            v = cell.value
            mark = ""
            if cell.coordinate in merged:
                rng = merged[cell.coordinate]
                mark = "⟨병합 시작⟩" if rng.startswith(cell.coordinate + ":") else "⟨병합⟩"
            s = "" if v is None else str(v)
            s = (s[:22] + "…") if len(s) > 22 else s
            cells.append((s + " " + mark).strip() or "·")
        print(f"| **{r}** | " + " | ".join(cells) + " |")
    if ws.max_row > max_rows:
        print(f"... {ws.max_row - max_rows}행 더 있음 (--max-rows 로 조절)")


def analyze(ws):
    merged = merged_lookup(ws)

    # 4) 헤더 단수 추정 — 첫 숫자 셀이 나오는 행 이전까지
    first_data_row = None
    for r in range(1, ws.max_row + 1):
        kinds = [cell_kind(ws.cell(row=r, column=c).value)
                 for c in range(1, ws.max_column + 1)]
        if kinds.count("숫자") + kinds.count("숫자형문자") >= 2:
            first_data_row = r
            break
    print("\n#### 헤더 단수")
    if first_data_row:
        print(f"- 첫 데이터 행: **{first_data_row}행** → 그 위 {first_data_row - 1}행이 제목·헤더 영역")
        for r in range(1, first_data_row):
            vals = [str(ws.cell(row=r, column=c).value)
                    for c in range(1, ws.max_column + 1)
                    if ws.cell(row=r, column=c).value is not None]
            has_merge = any(ws.cell(row=r, column=c).coordinate in merged
                            for c in range(1, ws.max_column + 1))
            print(f"  - {r}행: {len(vals)}개 값{' · 병합 있음' if has_merge else ''} — {vals[:6]}")
    else:
        print("- 숫자 행을 찾지 못했습니다")

    # 5) 열별 자료형 혼재 · 단위 문자열
    print("\n#### 열별 자료형 / 단위 혼재")
    print("| 열 | 헤더 추정 | 자료형 분포 | 단위 문자열 |")
    print("|:--|:--|:--|:--|")
    start = first_data_row or 1
    for c in range(1, ws.max_column + 1):
        kinds, units = {}, set()
        for r in range(start, ws.max_row + 1):
            v = ws.cell(row=r, column=c).value
            k = cell_kind(v)
            kinds[k] = kinds.get(k, 0) + 1
            if isinstance(v, str):
                units.update(RE_UNIT.findall(v))
        head = next((str(ws.cell(row=r, column=c).value)
                     for r in range(1, start) if ws.cell(row=r, column=c).value is not None), "")
        dist = ", ".join(f"{k} {n}" for k, n in sorted(kinds.items(), key=lambda kv: -kv[1])
                         if k != "빈칸")
        mixed = "**혼재**" if len([k for k in kinds if k not in ("빈칸",)]) > 1 else ""
        print(f"| {get_column_letter(c)} | {head[:20]} | {dist or '—'} {mixed} | "
              f"{', '.join(sorted(units)) if units else '—'} |")

    # 6) 각주 후보 — 아래쪽 연속 텍스트 행
    print("\n#### 각주 후보")
    found = []
    for r in range(ws.max_row, 0, -1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        nonnull = [v for v in vals if v is not None]
        if not nonnull:
            continue
        kinds = {cell_kind(v) for v in nonnull}
        if kinds <= {"문자"} and len(nonnull) <= 2:
            found.append((r, str(nonnull[0])))
        else:
            break
    for r, text in reversed(found):
        flag = " ← 각주 기호" if RE_FOOTNOTE.match(text) else ""
        print(f"- {r}행: {text}{flag}")
    if not found:
        print("- 없음")
    if found:
        print(f"\n→ 데이터 블록은 **{first_data_row}행 ~ {min(r for r, _ in found) - 1}행**")


def main(argv=None):
    ap = argparse.ArgumentParser(description="엑셀 구조 조사 (변환하지 않음)")
    ap.add_argument("path", nargs="?",
                    default=os.path.join(BASE, "exercises", "03", "internal_panel.xlsx"))
    ap.add_argument("--max-rows", type=int, default=30)
    ap.add_argument("--max-cols", type=int, default=12)
    a = ap.parse_args(argv)

    wb = load_workbook(a.path, data_only=True)
    print(f"파일: {os.path.relpath(a.path, BASE)}")
    print(f"시트 {len(wb.sheetnames)}개: {wb.sheetnames}")
    for name in wb.sheetnames:
        ws = wb[name]
        dump(ws, a.max_rows, a.max_cols)
        analyze(ws)
    return 0


if __name__ == "__main__":
    sys.exit(main())
