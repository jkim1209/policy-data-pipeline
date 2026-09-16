"""internal_panel.xlsx → 표준 스키마.

구조 조사 결과(inspect_xlsx.py)와 사용자가 확정한 해석을 그대로 옮깁니다.

확정된 해석 (2026-09-16)
  - 출생아수 단위는 명으로 통일. '43.5천' 같은 천 단위 표기는 환산
  - '-'  → NA_CONFIDENTIAL (비공개)      ← 각주 주2
  - 공란 → NA_NOTSURVEYED (미집계)       ← 각주 주2
  - 2025년은 vintage '잠정', 2023·2024는 '확정'  ← 각주 주3
  - 지역명이 references/region-codes.csv 에 없으면 중단
  - 각주 행(16~18)은 데이터로 읽지 않음

열 매핑 — 헤더가 데이터보다 한 칸 왼쪽으로 밀려 있습니다.
헤더 3행은 A~G 7열을 선언하지만 실제 데이터는 A~H 8열입니다. 권역 열이
헤더에 선언되지 않은 채 A 를 차지했기 때문입니다. 헤더를 그대로 믿으면
E열의 TFR 값이 '출생아수 2023'으로 들어갑니다. 아래는 조사로 확인한 실제 배치입니다.

    A 권역(병합)   B 시도명
    C TFR 2023     D TFR 2024     E TFR 2025
    F 출생아수2023 G 출생아수2024 H 출생아수2025
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from schema import COLUMNS, KST, Record, SchemaError, to_rows  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
XLSX = os.path.join(BASE, "exercises", "03", "internal_panel.xlsx")

SOURCE = "INTERNAL"
DATA_FIRST_ROW = 5          # 1 제목 / 2 공백 / 3 헤더1단 / 4 헤더2단 / 5~ 데이터
REGION_COL = 2              # B — 시도명
GROUP_COL = 1               # A — 권역(병합). 표준 스키마에 자리가 없어 싣지 않습니다

# (엑셀 열, indicator_code, 연도, 단위)
LAYOUT = [
    (3, "TFR", 2023, "명"), (4, "TFR", 2024, "명"), (5, "TFR", 2025, "명"),
    (6, "BIRTHS", 2023, "명"), (7, "BIRTHS", 2024, "명"), (8, "BIRTHS", 2025, "명"),
]

# 각주 주3 — 2025년만 잠정
VINTAGE_BY_YEAR = {2023: "확정", 2024: "확정", 2025: "잠정"}

# 각주 주2 — 표기별 사유 코드
MISSING_BY_MARK = {"-": "NA_CONFIDENTIAL", "": "NA_NOTSURVEYED"}

# 각주 주1 — 담당자가 천명 단위로 기재한 셀 (예: '43.5천')
RE_THOUSAND = re.compile(r"^\s*([\d.]+)\s*천\s*$")
RE_FOOTNOTE = re.compile(r"^\s*(주\d*\)|[*※注])")


class ConvertError(RuntimeError):
    pass


def region_map():
    path = os.path.join(BASE, "references", "region-codes.csv")
    with open(path, encoding="utf-8") as f:
        return {r["region_name"]: r["region_code"] for r in csv.DictReader(f)}


def parse_value(raw, indicator, coord, converted):
    """셀 값 → (실수 또는 None, 결측 사유 또는 None).

    환산은 여기 한 곳에서만 합니다. 환산한 셀은 전부 기록해 보고합니다.
    """
    if raw is None:
        return None, MISSING_BY_MARK[""]

    if isinstance(raw, str):
        s = raw.strip()
        if s in MISSING_BY_MARK:
            return None, MISSING_BY_MARK[s]
        m = RE_THOUSAND.match(s)
        if m:
            if indicator != "BIRTHS":
                raise ConvertError(
                    f"{coord}: 천 단위 표기 {s!r} 가 {indicator} 열에 있습니다. "
                    "각주 주1은 출생아수에만 해당합니다. 원본을 확인하세요."
                )
            value = float(m.group(1)) * 1000
            converted.append({"셀": coord, "원문": s, "환산": value, "근거": "각주 주1"})
            return value, None
        try:
            return float(s.replace(",", "")), None
        except ValueError:
            raise ConvertError(
                f"{coord}: 값을 해석할 수 없습니다 {raw!r}. "
                "결측 표기('-'·공란)도 천 단위 표기도 아닙니다."
            ) from None

    if isinstance(raw, (int, float)):
        return float(raw), None

    raise ConvertError(f"{coord}: 예상치 못한 자료형 {type(raw).__name__} ({raw!r})")


def convert(path=XLSX, log=print):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rmap = region_map()
    ts = datetime.now(KST).isoformat(timespec="seconds")
    url = f"file://{path}"

    records, converted, unmapped, skipped = [], [], [], []

    for r in range(DATA_FIRST_ROW, ws.max_row + 1):
        name = ws.cell(row=r, column=REGION_COL).value
        name = name.strip() if isinstance(name, str) else name

        # 각주 행은 데이터로 읽지 않습니다. 데이터 블록은 빈 행에서 끝납니다.
        if name in (None, ""):
            a = ws.cell(row=r, column=GROUP_COL).value
            if isinstance(a, str) and RE_FOOTNOTE.match(a):
                skipped.append({"행": r, "사유": "각주", "내용": a.strip()[:60]})
            else:
                skipped.append({"행": r, "사유": "빈 행", "내용": ""})
            continue

        if isinstance(name, str) and RE_FOOTNOTE.match(name):
            skipped.append({"행": r, "사유": "각주", "내용": name[:60]})
            continue

        code = rmap.get(name)
        if code is None:
            unmapped.append({"행": r, "지역명": name})
            continue

        for col, indicator, year, unit in LAYOUT:
            coord = f"{L(col)}{r}"
            value, reason = parse_value(ws.cell(row=r, column=col).value,
                                        indicator, coord, converted)
            try:
                records.append(Record(
                    source=SOURCE, indicator_code=indicator, region_code=code,
                    period=year, value=value, unit=unit,
                    vintage=VINTAGE_BY_YEAR[year], source_url=url,
                    retrieved_at=ts, missing_reason=reason,
                ))
            except SchemaError as e:
                raise ConvertError(f"{coord} ({name}): {e}") from None

    # 사전에 없는 지역명이 있으면 산출물을 만들지 않습니다.
    if unmapped:
        lines = "\n".join(f"    {u['행']}행: {u['지역명']!r}" for u in unmapped)
        raise ConvertError(
            f"references/region-codes.csv 에 없는 지역명 {len(unmapped)}건. 중단합니다.\n"
            f"{lines}\n"
            "    임의로 매핑하지 않습니다. 사전에 등록한 뒤 다시 실행하세요."
        )

    log(f"  데이터 행 {len(records) // len(LAYOUT)}개 지역 × {len(LAYOUT)}열 = {len(records)}행")
    for s in skipped:
        log(f"  건너뜀 {s['행']}행 [{s['사유']}] {s['내용']}")
    for c in converted:
        log(f"  환산 {c['셀']} {c['원문']} → {c['환산']:.0f} ({c['근거']})")
    return records, converted, skipped


def main(argv=None):
    ap = argparse.ArgumentParser(description="internal_panel.xlsx → 표준 스키마")
    ap.add_argument("--out", help="CSV 저장 경로")
    a = ap.parse_args(argv)

    try:
        records, converted, skipped = convert()
    except (ConvertError, SchemaError) as e:
        print(f"중단: {e}", file=sys.stderr)
        return 1

    rows = to_rows(records)          # 중복 키 검사
    try:
        import pandas as pd
        df = pd.DataFrame(rows, columns=COLUMNS)
        print(f"\n표준 스키마 {len(df)}행")
        print(f"\n지표·연도별 (결측 제외):\n"
              f"{df[df['value'] != ''].groupby(['indicator_code', 'period']).size().to_string()}")
        print(f"\nvintage:\n{df.groupby(['period', 'vintage']).size().to_string()}")
        miss = df[df['value'] == '']
        print(f"\n결측 {len(miss)}행:")
        print(miss[["indicator_code", "region_code", "period", "missing_reason"]].to_string(index=False))
        if a.out:
            df.to_csv(a.out, index=False, encoding="utf-8-sig")
            print(f"\n저장: {a.out}")
    except ImportError:
        for r in rows:
            print(r)
    return 0


if __name__ == "__main__":
    sys.exit(main())
