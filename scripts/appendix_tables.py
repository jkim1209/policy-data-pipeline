"""보고서 부록용 표 생성.

out/panel.csv 에서 두 표를 만듭니다.
    부록 표 1. 시도별 합계출산율
    부록 표 2. 합계출산율 국제비교 (OECD 주요국)

출처 · 기준시점 · 단위는 **패널과 config/sources.yml 에서 자동으로** 뽑습니다.
표 아래에 손으로 적으면 갱신할 때마다 어긋납니다. 각주가 본문보다 먼저
낡습니다. 그래서 여기서는 각주를 데이터에서 만들어 냅니다.

각주에 들어가는 것
    출처     config 의 소스명 · 기관 · 통계표 ID · 요청 주소
    기준시점 vintage · 수록 기간 · 수집 시각 · (있으면) 원자료 최종 갱신일
    단위     패널의 unit 컬럼. 표 안에서 단위가 섞이면 중단합니다
    결측     실제로 나타난 사유 코드만. 없으면 각주도 없습니다
"""
from __future__ import annotations

import argparse
import csv
import glob
import json
import os
import sys
from datetime import datetime

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


class AppendixError(RuntimeError):
    pass


MISSING_LABEL = {
    "NA_NOTSURVEYED": "미조사",
    "NA_NOTAPPLICABLE": "해당없음",
    "NA_CONFIDENTIAL": "비공개",
}


def load_config():
    import yaml
    with open(os.path.join(BASE, "config", "sources.yml"), encoding="utf-8") as f:
        return yaml.safe_load(f)


def load_dict(fname, key, val):
    path = os.path.join(BASE, "references", fname)
    with open(path, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    return {r[key]: r[val] for r in rows}, [r[key] for r in rows]


def load_panel(path):
    import pandas as pd
    df = pd.read_csv(path, encoding="utf-8-sig")
    df["period"] = df["period"].astype(int)
    return df


def source_entry(config, indicator_code):
    """패널의 indicator_code 로 config 소스 항목을 되찾습니다."""
    for s in config["sources"]:
        if s.get("indicator_code") == indicator_code:
            return s
    raise AppendixError(
        f"config/sources.yml 에 indicator_code={indicator_code!r} 인 소스가 없습니다. "
        "각주의 출처를 지어내지 않습니다."
    )


def raw_lastupdated(source_id):
    """원자료의 최종 갱신일. raw/<실행일>/<소스>.json 에서 찾습니다.

    World Bank 는 응답 메타에 lastupdated 를 줍니다. 수집 시각과 다른 값이고,
    보고서에서 '기준시점'으로 인용해야 하는 쪽은 이것입니다.
    """
    # raw/sample/ 은 수업용 픽스처입니다. 실행 스냅샷(raw/<실행일>/)만 봅니다.
    paths = sorted(
        p for p in glob.glob(os.path.join(BASE, "raw", "*", f"{source_id}.json"))
        if os.path.basename(os.path.dirname(p)) != "sample"
    )
    if not paths:
        return None
    try:
        with open(paths[-1], encoding="utf-8") as f:
            payload = json.load(f)
    except (OSError, ValueError):
        return None
    if isinstance(payload, list) and len(payload) == 2 and isinstance(payload[0], dict):
        return payload[0].get("lastupdated")
    return None


def fmt(v, decimals):
    if v is None or v != v:
        return None
    return f"{v:,.{decimals}f}"


def build_table(df, indicator, order, labels, config, title, decimals,
                order_note, head="지역", extra_note=None):
    sel = df[df["indicator_code"] == indicator].copy()
    if sel.empty:
        raise AppendixError(f"패널에 {indicator} 행이 없습니다")

    units = sorted(set(sel["unit"]))
    if len(units) != 1:
        raise AppendixError(
            f"{indicator}: 표 안에 단위가 섞여 있습니다 {units}. "
            "부록 표는 단위 하나로만 만듭니다."
        )
    unit = units[0]

    years = sorted(set(sel["period"]))
    cell = {(r.region_code, r.period): r for r in sel.itertuples()}
    codes = [c for c in order if any((c, y) in cell for y in years)]
    missing_seen, rows, blanks = set(), [], 0

    for code in codes:
        cells = []
        for y in years:
            r = cell.get((code, y))
            if r is None:
                cells.append("")
                blanks += 1
                continue
            s = fmt(r.value, decimals)
            if s is None:
                reason = r.missing_reason if isinstance(r.missing_reason, str) else ""
                missing_seen.add(reason)
                cells.append(f"–{MISSING_LABEL.get(reason, '결측')}")
            else:
                cells.append(s)
        rows.append((labels.get(code, code), code, cells))

    src = source_entry(config, indicator)
    vintages = sorted(set(sel["vintage"]))
    retrieved = sorted(set(sel["retrieved_at"]))[-1]
    lastupd = raw_lastupdated(src["id"])
    urls = sorted(set(sel["source_url"]))

    notes = []
    org = "통계청" if src["id"].startswith("kosis") else (
        "World Bank" if src["id"].startswith("wb") else "")
    tbl = src.get("params", {}).get("tblId")
    itm = src.get("params", {}).get("itmId")
    origin = f"{org} · " if org else ""
    detail = f" (통계표 {tbl}, 항목 {itm})" if tbl else f" (지표 {indicator})"
    notes.append(f"**출처**: {origin}{src['name']}{detail}. 요청 주소 {urls[0]}")

    ts = datetime.fromisoformat(retrieved).strftime("%Y-%m-%d %H:%M")
    base_time = f"자료 {years[0]}~{years[-1]}년, {'·'.join(vintages)}"
    if lastupd:
        base_time += f" · 원자료 최종 갱신 {lastupd}"
    notes.append(f"**기준시점**: {base_time} · 수집 {ts} (KST)")
    notes.append(f"**단위**: {unit}")
    if missing_seen:
        seen = ", ".join(f"–{MISSING_LABEL.get(m, m)}(`{m}`)" for m in sorted(missing_seen))
        tail = " 빈칸은 해당 연도가 수집 범위에 없음을 뜻합니다." if blanks else ""
        notes.append(f"**결측**: {seen}.{tail}")
    notes.append(f"**정렬**: {order_note}")
    if extra_note:
        notes.append(extra_note)

    return {"title": title, "years": years, "rows": rows, "unit": unit,
            "notes": notes, "head": head, "blanks": blanks}


def to_markdown(t, number):
    out = [f"## 부록 표 {number}. {t['title']}", ""]
    out.append(f"| {t['head']} | " + " | ".join(str(y) for y in t["years"]) + " |")
    out.append("|:--|" + "---:|" * len(t["years"]))
    for label, _code, cells in t["rows"]:
        out.append(f"| {label} | " + " | ".join(cells) + " |")
    out.append("")
    for n in t["notes"]:
        out.append(f"- {n}")
    out.append("")
    return "\n".join(out)


def to_sheet(wb, t, number):
    from openpyxl.styles import Alignment, Font
    ws = wb.create_sheet(f"부록표{number}")
    ws.append([f"부록 표 {number}. {t['title']}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append([t["head"]] + [str(y) for y in t["years"]])
    for c in ws[3]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    for label, _code, cells in t["rows"]:
        row = [label]
        for v in cells:
            try:
                row.append(float(v.replace(",", "")))
            except (ValueError, AttributeError):
                row.append(v)
        ws.append(row)
    ws.append([])
    for n in t["notes"]:
        ws.append([n.replace("**", "")])
    ws.column_dimensions["A"].width = 18
    for col in ws.iter_cols(min_col=2, max_col=len(t["years"]) + 1):
        ws.column_dimensions[col[0].column_letter].width = 9
    return ws


def main(argv=None):
    ap = argparse.ArgumentParser(description="보고서 부록용 표 생성")
    ap.add_argument("--panel", default=os.path.join(BASE, "out", "panel.csv"))
    ap.add_argument("--out-dir", default=os.path.join(BASE, "out", "appendix"))
    a = ap.parse_args(argv)

    config = load_config()
    df = load_panel(a.panel)

    region_names, region_order = load_dict("region-codes.csv", "region_code", "region_name")
    country_names, country_order = load_dict("country-names.csv", "country_code", "country_name")

    # 국제비교는 한국을 맨 위에, 나머지는 최근 연도 값이 큰 순서로 둡니다.
    wb_rows = df[df["indicator_code"] == "SP.DYN.TFRT.IN"]
    # 최신 연도는 전 국가가 비어 있을 수 있습니다(공표 전). 값이 실제로 있는
    # 가장 최근 연도를 정렬 기준으로 삼고, 각주에도 그 연도를 적습니다.
    filled = wb_rows[wb_rows["value"].notna()]
    if filled.empty:
        raise AppendixError("국제비교 표에 값이 있는 행이 없습니다")
    latest = int(filled["period"].max())
    have = wb_rows[(wb_rows["period"] == latest) & wb_rows["value"].notna()]
    ranked = list(have.sort_values("value", ascending=False)["region_code"])
    intl_order = ["KOR"] + [c for c in ranked if c != "KOR"] + \
                 [c for c in country_order if c not in ranked and c != "KOR"]

    t1 = build_table(
        df, "TFR", region_order, region_names, config,
        "시도별 합계출산율", 3,
        "`references/region-codes.csv` 의 행정구역 코드 순. 전국이 첫 행입니다.",
        head="지역",
    )
    t2 = build_table(
        df, "SP.DYN.TFRT.IN", intl_order, country_names, config,
        "합계출산율 국제비교 (OECD 주요국)", 3,
        f"대한민국을 첫 행에 두고, 나머지는 값이 있는 최신 연도({latest}년) 기준 내림차순입니다.",
        head="국가",
        extra_note=(
            "**주의**: 이 표의 출처는 **World Bank Indicators** 이며 OECD 통계가 아닙니다. "
            "표에 실린 7개국이 모두 OECD 회원국이라 '국제비교'로 묶은 것입니다. "
            "OECD 공표치를 인용해야 한다면 `references/api-registry.md` 의 OECD 항목을 "
            "먼저 검증해야 합니다(현재 미검증)."
        ),
    )

    os.makedirs(a.out_dir, exist_ok=True)
    md = ["# 부록 — 통계표", "",
          f"생성 {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M')} (KST) · "
          f"원자료 `out/panel.csv`", ""]
    md.append(to_markdown(t1, 1))
    md.append(to_markdown(t2, 2))
    md_path = os.path.join(a.out_dir, "appendix.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md))

    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    to_sheet(wb, t1, 1)
    to_sheet(wb, t2, 2)
    xlsx_path = os.path.join(a.out_dir, "appendix.xlsx")
    wb.save(xlsx_path)

    print(f"부록 표 2종 생성")
    print(f"  {md_path}")
    print(f"  {xlsx_path}")
    print(f"\n표 1: {len(t1['rows'])}개 지역 × {len(t1['years'])}년 · 단위 {t1['unit']}")
    print(f"표 2: {len(t2['rows'])}개국 × {len(t2['years'])}년 · 단위 {t2['unit']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
